# -*- coding: utf-8 -*-
"""从《打印价格暂定方案和材料公示.xlsx》读取默认数据。"""

import re

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

from . import defaults as D
from .defaults import _split_sides


class ExcelImportError(Exception):
    pass


_PRICE_RE = re.compile(r"([0-9]+(?:\.[0-9]+)?)")
_LIMIT_RE = re.compile(r"[≤<]\s*([0-9]+(?:\.[0-9]+)?)\s*%")


def _text(value):
    if value is None:
        return ""
    return str(value).strip()


def _iter_rows(ws):
    for row in ws.iter_rows(values_only=True):
        yield [_text(c) for c in row]


def _find_header(rows, keyword):
    for index, row in enumerate(rows):
        if any(cell == keyword for cell in row):
            return index
    return None


def _col(row, index):
    if index is None or index >= len(row):
        return ""
    return row[index]


def _cell_index(header, *names):
    for name in names:
        for idx, cell in enumerate(header):
            if cell == name:
                return idx
    return None


def _parse_price(text):
    text = _text(text)
    if not text or "暂不" in text or "不提供" in text:
        return None, False
    match = _PRICE_RE.search(text)
    if not match:
        return None, False
    return float(match.group(1)), True


def _parse_limit(text):
    match = _LIMIT_RE.search(_text(text))
    if not match:
        return None
    value = float(match.group(1))
    return value / 100.0 if value > 1 else value


def _sheet_lines(wb):
    return {ws.title: list(_iter_rows(ws)) for ws in wb.worksheets}


def _parse_items(lines):
    for rows in lines.values():
        header_idx = _find_header(rows, "收费标准")
        if header_idx is None:
            continue
        header = rows[header_idx]
        if "纸张介质" not in header:
            continue
        c_no = _cell_index(header, "序号")
        c_paper = _cell_index(header, "纸张介质")
        c_spec = _cell_index(header, "规格尺寸 / 克重", "规格尺寸/克重")
        c_method = _cell_index(header, "打印方式")
        c_color = _cell_index(header, "色彩模式")
        c_sides = _cell_index(header, "打印面数")
        c_price = _cell_index(header, "收费标准")
        c_note = _cell_index(header, "计费说明与适用场景")

        items = []
        cur_paper = cur_spec = cur_method = cur_color = ""
        cur_limit = None
        cur_note = ""
        seq = 0
        for row in rows[header_idx + 1:]:
            paper = _col(row, c_paper) or cur_paper
            if _col(row, c_paper):
                cur_limit = None
                cur_note = ""
            spec = _col(row, c_spec) or cur_spec
            method = _col(row, c_method) or cur_method
            color = _col(row, c_color) or cur_color
            sides_text = _col(row, c_sides)
            price_text = _col(row, c_price)
            raw_note = _col(row, c_note)
            note = raw_note or cur_note
            if raw_note:
                cur_note = raw_note

            cur_paper, cur_spec, cur_method, cur_color = paper, spec, method, color

            if not (paper and sides_text and price_text):
                continue

            price, available = _parse_price(price_text)
            limit = _parse_limit(note)
            if limit is not None:
                cur_limit = limit
            if limit is None:
                limit = cur_limit
            if price is None:
                limit = None

            seq += 1
            items.append({
                "id": seq,
                "paper": paper,
                "spec": spec,
                "method": method,
                "color": color,
                "sides": _split_sides(sides_text),
                "price": price,
                "limit": limit,
                "note": note,
                "available": available,
            })
        if items:
            return items
    return []


def _parse_materials(lines):
    for rows in lines.values():
        header_idx = _find_header(rows, "耗材分类")
        if header_idx is None:
            continue
        header = rows[header_idx]
        idx = [_cell_index(header, n) for n in
               ("序号", "耗材分类", "物品名称", "品牌 / 厂商", "品牌/厂商",
                "规格型号 / 核心参数", "规格型号/核心参数", "包装规格", "主要用途 / 适用场景", "主要用途/适用场景")]
        c_no, c_cat, c_name, c_brand1, c_brand2, c_spec1, c_spec2, c_pack, c_use1, c_use2 = idx
        c_brand = c_brand1 if c_brand1 is not None else c_brand2
        c_spec = c_spec1 if c_spec1 is not None else c_spec2
        c_use = c_use1 if c_use1 is not None else c_use2
        out = []
        cur_cat = ""
        for row in rows[header_idx + 1:]:
            name = _col(row, c_name)
            if not name:
                continue
            if _col(row, c_cat):
                cur_cat = _col(row, c_cat)
            out.append({
                "no": _col(row, c_no),
                "category": cur_cat,
                "name": name,
                "brand": _col(row, c_brand),
                "spec": _col(row, c_spec),
                "pack": _col(row, c_pack),
                "usage": _col(row, c_use),
            })
        if out:
            return out
    return []


def _parse_equipment(lines):
    for rows in lines.values():
        header_idx = _find_header(rows, "资产类别")
        if header_idx is None:
            continue
        header = rows[header_idx]
        c_no = _cell_index(header, "序号")
        c_cat = _cell_index(header, "资产类别")
        c_brand = _cell_index(header, "厂商 / 品牌", "厂商/品牌")
        c_name = _cell_index(header, "设备 / 软件名称", "设备/软件名称")
        c_spec = _cell_index(header, "版本型号 / 规格参数", "版本型号/规格参数")
        out = []
        cur_cat = ""
        for row in rows[header_idx + 1:]:
            name = _col(row, c_name)
            if not name:
                continue
            if _col(row, c_cat):
                cur_cat = _col(row, c_cat)
            out.append({
                "no": _col(row, c_no),
                "category": cur_cat,
                "brand": _col(row, c_brand),
                "name": name,
                "spec": _col(row, c_spec),
            })
        if out:
            return out
    return []


def _parse_notes(lines):
    notes = []
    for rows in lines.values():
        for row in rows:
            first = row[0] if row else ""
            if re.match(r"^\d+[.、]\s*\S", first) and len(first) > 8:
                if first not in notes:
                    notes.append(first)
    return notes


def import_from_excel(path):
    """读取 Excel，返回可合并进配置的字典。"""
    if openpyxl is None:
        raise ExcelImportError("未安装 openpyxl，无法读取 Excel")
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ExcelImportError("无法打开 Excel：%s" % exc) from exc

    lines = _sheet_lines(wb)
    data = {}
    items = _parse_items(lines)
    if items:
        existing = {it["paper"] for it in items}
        for extra in D.DEFAULT_ITEMS:
            if extra["paper"] not in existing:
                items.append(dict(extra))
        for idx, item in enumerate(items, start=1):
            item["id"] = idx
        data["items"] = items

    materials = _parse_materials(lines)
    if materials:
        data["materials"] = materials
    equipment = _parse_equipment(lines)
    if equipment:
        data["equipment"] = equipment
    notes = _parse_notes(lines)
    if notes:
        data["notes"] = notes
    return data
